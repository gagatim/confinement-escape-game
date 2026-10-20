import docx,copy,re,json
SRC="C:/Users/Lenovo/Desktop/禁閉逃殺_職業分析報告_含模擬結果_更新.docx"
d=docx.Document(SRC)
t=d.tables[-1]
old={}
for r in t.rows[1:]:
    c=[x.text for x in r.cells]; old[c[0].split(' ',1)[1]]=(c[0],c[8])
lines=open("C:/Users/Lenovo/Desktop/禁閉逃殺/tools/analysis_new.txt",encoding="utf-8").read().splitlines()
head=lines[0]; rows=[l.split("|") for l in lines[2:] if l.strip()]
m=re.search(r"全體平均存活輪數 ([\d.]+)（標準差 ([\d.]+)）｜存活率 ([\d.]+)%",head); gm,gsd,sr=m.groups()
def st(cell,text):
    p=cell.paragraphs[0] if hasattr(cell,'paragraphs') else cell
    rs=p.runs
    if rs:
        rs[0].text=text
        for r in rs[1:]: r._r.getparent().remove(r._r)
    else: p.add_run(text)
def judge(name,z,th):
    z=float(z); th=float(th) if th else None
    if z>=2: s="模擬明顯偏強"; 
    elif z<=-2: s="模擬明顯偏弱"
    else: s="接近平均"
    if th is not None:
        if z>=2 and th>=7.5: s+="；與理論評估一致"
        elif z<=-2 and th<=4.5: s+="；與理論評估一致"
        elif z<=-2 and th>=7: s+="；與理論評估相反（AI 可能不會運用）"
        elif -2<z<2 and th>=8: s+="；低於理論評估（AI 可能不會運用）"
    return s
while len(t.rows)-1<len(rows): t._tbl.append(copy.deepcopy(t.rows[-1]._tr))
while len(t.rows)-1>len(rows): t._tbl.remove(t.rows[-1]._tr)
for r,(n,g,m_,z,alive,win,k,dl,tk) in zip(t.rows[1:],rows):
    disp,th=old.get(n,("　"+n,""))
    vals=[disp,g,m_,z,alive,k,dl,tk,th or "—",judge(n,z,th)]
    for c,v in zip(r.cells,vals): st(c,v)
S={r[0]:(float(r[2]),float(r[3])) for r in rows}
hi=[f"{n}（z {z}）" for n,(m_,z) in S.items() if z>=2]
lo=[f"{n}（{z}）" for n,(m_,z) in S.items() if z<=-2]
P=d.paragraphs
def sp(i,text): st(P[i],text)
sp(46,"方法：用電腦把 8 位玩家全部交給遊戲內建的 AI，自動打完 6 輪（擴充職業全開、一般模式、8 人局），共 700 局；每位玩家的職業隨機。每個職業出場約 90~130 次。本次使用「2026/10/5 調整後的最新規則」重新模擬（守護者、攀岩者、馴獸師、走私客等已修改）。模擬程式與原始紀錄放在專案的 tools 資料夾，可重跑。")
sp(47,f"全體平均存活 {gm} 輪（活到最後算 7 輪，標準差 {gsd}），僅 {sr}% 的玩家活到第 6 輪結束。主要指標是「存活輪數」；z 值＝（職業平均－全體平均）÷標準誤，|z|≥2 視為有明顯差異。")
g=lambda n:S[n][0]
sp(49,"明顯偏強："+"、".join(hi)+f"。上一輪最強的守護者（5.63）與攀岩者（5.55）修正後已回到平均附近（現為 {g('守護者'):.2f}、{g('攀岩者'):.2f}）；馴獸師改為「行動結束後才能使用巡迴獵犬」後由 4.78 降到 {g('馴獸師'):.2f}；黑社會老大由 5.06 降到 {g('黑社會老大'):.2f}。")
sp(50,"明顯偏弱："+"、".join(lo)+f"。毒人仍是最大離群值（{g('毒人'):.2f}，與其他職業差距極大），原因多半是 AI 不會刻意走進已中毒的功能房間，需要真人實測再判斷是否真的過弱。賭徒（{g('賭徒'):.2f}）與購物大媽（{g('購物大媽'):.2f}）在這次模擬中接近平均或略低，與理論評估的「偏弱」大致相符。")
sp(51,f"與理論評估不符：礦工（{g('礦工'):.2f}）、醫官（{g('醫官'):.2f}）、廚師（{g('廚師'):.2f}）、雇傭兵（{g('雇傭兵'):.2f}）理論上很強，但模擬中只是平均。主因是 AI 不會把金條換成水糧（礦工）、不會固定待在手術室（醫官）、不會拿刀當槍壓制（雇傭兵）。這類職業的實際強度取決於真人怎麼玩，不能用模擬下結論。")
sp(52,f"新職業：縱火犯 {g('縱火犯'):.2f}、吸血鬼 {g('吸血鬼'):.2f}、刺客 {g('刺客'):.2f}（刺客仍略低，AI 不會刻意尋找單人房間）。走私客（加上武力上限5後）{g('走私客'):.2f}，在平均附近。")
sp(56,f"守護者、攀岩者、馴獸師、走私客的調整都已生效，存活輪數已回到平均附近，過強問題已解決。接下來可處理礦工、醫官、雇傭兵（理論上的S級，模擬未證實，要靠真人實測判斷），以及庸才使用垃圾換血的漏洞。")
sp(57,f"目前模擬偏強的是超度師、生存狂、肥宅、邊緣人、求生教官，可再觀察；偏弱的是毒人（離群值，需確認是 AI 問題還是機制問題）、草藥師、工程師。理論上偏弱的賭徒、購物大媽、暗影使者仍建議小幅加強。")
d.save(SRC)
print("saved",len(rows),hi,lo)
